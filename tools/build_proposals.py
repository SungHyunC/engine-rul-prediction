"""Build bilingual one-page proposals with the Codex bundled Python runtime.

The PDF and editable DOCX share the same content. PDFs are independently typeset
with ReportLab; this does not certify DOCX pagination in Microsoft Word.
"""
from __future__ import annotations

import argparse
import subprocess
import sys
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
NASA = "https://www.nasa.gov/intelligent-systems-division/discovery-and-systems-health/pcoe/pcoe-data-set-repository/"
PAPER = "https://doi.org/10.3390/data6010005"

CONTENT = {
    "ko": {
        "title": "항공기 엔진 잔여수명 예측",
        "subtitle": "대용량 센서 데이터 병렬 처리 프로젝트 제안서",
        "meta": "김성현  |  학번 2021271250  |  제안서 제출 2026년 10월 21일",
        "sections": [
            ("연구 목표", "NASA N-CMAPSS DS02의 센서 시계열로 고장까지 남은 비행 사이클 수인 잔여수명(RUL)을 예측한다. 예측 정확도와 함께 데이터 읽기·특징 추출·병렬 처리 비용을 측정하고, 통계 특징 기반 회귀와 소형 1D CNN을 비교하는 재현 가능한 분석·추론 파이프라인을 구현한다."),
            ("데이터와 예측 방법", "DS02는 모의 엔진 9대의 6,517,190행이다[1, 2]. 운항 조건 W와 측정 센서 Xs의 18개 채널만 입력으로 사용하며, 열화 변수 T·건강 상태 hs·가상 센서 Xv는 제외한다. 엔진·비행 사이클 경계를 보존한 윈도에서 평균·표준편차·최솟값·최댓값·기울기의 90개 특징을 만든다. 평균 기준모델·Ridge·히스토그램 그래디언트 부스팅과 원시 시계열 1D CNN을 비교한다. 6가지 윈도·간격 설정과 원시/비음수 예측 정책도 평가한다."),
            ("평가와 빅데이터 실험", "개발 엔진 2, 5, 10, 16, 18, 20의 엔진별 3분할 교차검증으로 설정을 선택한다. 윈도 예측을 비행 종료 시점에 평균하고 엔진별 사이클 MAE·RMSE와 엔진 간 단순평균을 보고한다. 시험 엔진 11, 14, 15는 최초 실험에서 이미 열람했으므로 확장 결과는 동일 시험 자료의 탐색적 재평가로 명시한다. 작업자 1·2·4·8개와 HDF5 청크 크기를 바꾸어 시간·처리량·메모리·결과 일치를 측정한다. 엔진 3대의 시험 결과를 실제 항공기 운용 성능으로 일반화하지 않는다."),
        ],
        "schedule_heading": "구현 기록과 제출 일정",
        "schedule_header": ("기간", "수행 내용과 제출물"),
        "schedule": [
            ("10월 2일", "데이터·파이프라인·기준모델·민감도·CNN·추론 구현 및 실험"),
            ("10월 21일", "한국어 및 영어 각 1쪽 제안서 제출"),
            ("12월 8일", "5분 발표"),
            ("12월 16일 23:59", "보고서와 소스 코드 제출"),
        ],
        "refs_heading": "참고문헌",
        "source_note": "원출처는 NASA이며 실제 DS02 수집은 Figshare 제3자 미러(doi:10.6084/m9.figshare.20436504.v1)를 사용했다. NASA 전체 묶음 15.76 GB와 DS02 파일 2.45 GB를 구별한다.",
    },
    "en": {
        "title": "Aircraft Engine Remaining Useful Life Prediction",
        "subtitle": "Project Proposal for Parallel Processing of Large Sensor Data",
        "meta": "김성현  |  Student ID 2021271250  |  Proposal due 21 October 2026",
        "sections": [
            ("Research objective", "Predict remaining useful life (RUL), measured in flight cycles until failure, from NASA N-CMAPSS DS02. Build a reproducible analysis and inference pipeline that compares regression on statistical features with a compact 1D CNN, while measuring reading, feature-extraction and parallel-processing costs."),
            ("Data and prediction methods", "DS02 has 6,517,190 rows from nine simulated engines [1, 2]. Inputs contain only 18 operating-condition and measured-sensor channels (W and Xs), excluding degradation variables T, health state hs and virtual sensors Xv. Engine and flight-cycle boundaries are preserved. Mean, standard deviation, minimum, maximum and slope provide 90 features for a mean baseline, Ridge and histogram gradient boosting. A 1D CNN uses raw sensor windows. Six window/stride settings and raw/nonnegative prediction policies are also evaluated."),
            ("Evaluation and computing experiments", "Select settings using three-fold engine-grouped validation on development engines 2, 5, 10, 16, 18 and 20. Average window predictions at flight end, then report cycle-level MAE and RMSE per engine and their unweighted means. Test engines 11, 14 and 15 were inspected in the original experiment; extensions are explicitly exploratory re-evaluations of this same test split. Compare 1, 2, 4 and 8 workers and HDF5 chunk sizes for time, throughput, memory and output agreement. Three simulated test engines do not establish real-aircraft operational performance."),
        ],
        "schedule_heading": "Implementation record and submission dates",
        "schedule_header": ("Dates in 2026", "Work and deliverables"),
        "schedule": [
            ("2 October", "Data, pipeline, baselines, sensitivity study, CNN, inference and experiments"),
            ("21 October", "Submit separate one-page Korean and English proposals"),
            ("8 December", "Five-minute presentation"),
            ("16 December 23:59", "Submit report and source code"),
        ],
        "refs_heading": "References",
        "source_note": "NASA is the original source; DS02 was acquired through the Figshare third-party mirror (doi:10.6084/m9.figshare.20436504.v1). The NASA full bundle is 15.76 GB; the DS02 file is 2.45 GB.",
    },
}


def set_font(style, name="Malgun Gothic", size=11, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), "Malgun Gothic")
    rpr.append(rfonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    rpr.append(color)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def write_docx(lang, target):
    c = CONTENT[lang]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2), Cm(2)
    for name in ("Normal", "Title", "Subtitle", "Heading 1"):
        set_font(doc.styles[name])
    normal = doc.styles["Normal"]
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(5)
    set_font(doc.styles["Title"], size=17, bold=True)
    doc.styles["Title"].paragraph_format.space_after = Pt(4)
    set_font(doc.styles["Subtitle"], size=11, bold=True)
    doc.styles["Subtitle"].paragraph_format.space_after = Pt(6)
    set_font(doc.styles["Heading 1"], size=11.5, bold=True)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(8)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(3)
    doc.add_paragraph(c["title"], "Title")
    doc.add_paragraph(c["subtitle"], "Subtitle")
    p = doc.add_paragraph(c["meta"])
    for run in p.runs:
        run.font.size = Pt(9.5)
    for heading, text in c["sections"]:
        doc.add_paragraph(heading, "Heading 1")
        doc.add_paragraph(text)
    doc.add_paragraph(c["schedule_heading"], "Heading 1")
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Cm(4.2), Cm(12.8))
    for cell, value, width in zip(table.rows[0].cells, c["schedule_header"], widths):
        cell.text = value
        cell.width = width
    for row in c["schedule"]:
        cells = table.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.text = value
            cell.width = width
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement(f"w:{edge}")
                for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    border.set(qn(f"w:{key}"), value)
                borders.append(border)
            tcpr.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge, value in (("top", "50"), ("bottom", "50"), ("left", "80"), ("right", "80")):
                node = OxmlElement(f"w:{edge}")
                node.set(qn("w:w"), value)
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            tcpr.append(margins)
            if i == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F0F0F0")
                tcpr.append(shading)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.bold = i == 0
    doc.add_paragraph(c["refs_heading"], "Heading 1")
    p = doc.add_paragraph()
    add_hyperlink(p, "[1] NASA PCoE Data Set Repository, Dataset 17, Turbofan Engine Degradation Simulation-2.", NASA)
    p2 = doc.add_paragraph()
    add_hyperlink(p2, "[2] Arias Chao et al. (2021). Aircraft Engine Run-to-Failure Dataset under Real Flight Conditions for Prognostics and Diagnostics. Data, 6(1), 5. doi:10.3390/data6010005.", PAPER)
    p3 = doc.add_paragraph(c["source_note"])
    for p in (p, p2, p3):
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            run.font.size = Pt(8)
        for runpr in p._p.findall('.//' + qn('w:rPr')):
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), '16')
            runpr.append(sz)
    doc.core_properties.title = c["title"]
    doc.core_properties.author = "김성현"
    doc.core_properties.subject = "N-CMAPSS DS02 remaining useful life project proposal"
    doc.save(target)


def write_pdf(lang, target, font_dir):
    c = CONTENT[lang]
    regular_font = f"Malgun-{lang}"
    bold_font = f"Malgun-{lang}-Bold"
    pdfmetrics.registerFont(TTFont(regular_font, str(font_dir / "malgun.ttf")))
    pdfmetrics.registerFont(TTFont(bold_font, str(font_dir / "malgunbd.ttf")))
    pdfmetrics.registerFontFamily(regular_font, normal=regular_font, bold=bold_font,
                                 italic=regular_font, boldItalic=bold_font)
    wrap = "CJK" if lang == "ko" else None
    body = ParagraphStyle("Body", fontName=regular_font, fontSize=11.2, leading=15.2, spaceAfter=4,
                          textColor=colors.black, wordWrap=wrap, alignment=TA_LEFT)
    title = ParagraphStyle("Title", parent=body, fontName=bold_font, fontSize=17,
                           leading=21, spaceAfter=4)
    subtitle = ParagraphStyle("Subtitle", parent=body, fontName=bold_font, fontSize=10.7,
                              leading=14, spaceAfter=6)
    meta = ParagraphStyle("Meta", parent=body, fontSize=9, leading=12, spaceAfter=4, wordWrap="CJK")
    heading = ParagraphStyle("Heading", parent=body, fontName=bold_font, fontSize=11.4,
                             leading=14, spaceBefore=7, spaceAfter=3, keepWithNext=True)
    small = ParagraphStyle("Small", parent=body, fontSize=8, leading=10.3, spaceAfter=2)
    cell = ParagraphStyle("Cell", parent=body, fontSize=9, leading=11.8, spaceAfter=0)
    cellhead = ParagraphStyle("CellHead", parent=cell, fontName=bold_font)
    story = [Paragraph(escape(c["title"]), title), Paragraph(escape(c["subtitle"]), subtitle),
             Paragraph(escape(c["meta"]), meta)]
    for name, text in c["sections"]:
        story.extend([Paragraph(escape(name), heading), Paragraph(escape(text), body)])
    story.append(Paragraph(c["schedule_heading"], heading))
    rows = [[Paragraph(escape(x), cellhead) for x in c["schedule_header"]]]
    rows.extend([[Paragraph(escape(x), cell) for x in row] for row in c["schedule"]])
    table = Table(rows, colWidths=[42*mm, 128*mm], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("GRID", (0,0), (-1,-1), 0.45, colors.HexColor("#D9D9D9")),
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#F0F0F0")),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("LEFTPADDING", (0,0), (-1,-1), 5), ("RIGHTPADDING", (0,0), (-1,-1), 5),
        ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]))
    story.append(table)
    story.append(Paragraph(c["refs_heading"], heading))
    story.append(Paragraph(f'<link href="{NASA}" color="black">[1] NASA PCoE Data Set Repository, Dataset 17, Turbofan Engine Degradation Simulation-2.</link>', small))
    story.append(Paragraph(f'<link href="{PAPER}" color="black">[2] Arias Chao et al. (2021). Aircraft Engine Run-to-Failure Dataset under Real Flight Conditions for Prognostics and Diagnostics. <i>Data</i>, 6(1), 5. doi:10.3390/data6010005.</link>', small))
    story.append(Paragraph(escape(c["source_note"]), small))
    doc = SimpleDocTemplate(str(target), pagesize=A4, leftMargin=20*mm, rightMargin=20*mm,
                            topMargin=18*mm, bottomMargin=18*mm, title=c["title"], author="김성현")
    doc.build(story)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=ROOT / "output")
    parser.add_argument("--font-dir", type=Path, default=Path("C:/Windows/Fonts"))
    parser.add_argument("--lang", choices=("ko", "en"))
    args = parser.parse_args()
    # Isolate ReportLab font subsetting state for the two language artifacts.
    if args.lang is None:
        for lang in ("ko", "en"):
            subprocess.run([sys.executable, str(Path(__file__).resolve()), "--lang", lang,
                            "--output-dir", str(args.output_dir), "--font-dir", str(args.font_dir)], check=True)
        return
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for lang in (args.lang,):
        for extension, builder in (("docx", write_docx), ("pdf", write_pdf)):
            target = args.output_dir / f"proposal_{lang}.{extension}"
            if extension == "pdf":
                builder(lang, target, args.font_dir)
            else:
                builder(lang, target)
            print(target)


if __name__ == "__main__":
    main()
